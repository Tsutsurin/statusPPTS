"""Общие функции чтения таблиц: строки листов, поиск столбцов по заголовкам, разбор ячеек."""

from __future__ import annotations

import csv
import logging
import re
from pathlib import Path

import openpyxl

from status.core.models import PptsColumnMapping

logger = logging.getLogger(__name__)


class ReaderError(Exception):
    ...


PPTS_EXACT_ID = "ID ПТС"
PPTS_EXACT_NAME = "Название ПТС"
PPTS_EXACT_VENDOR = "Вендор"
PPTS_NR_MARK = "Nr."
PPTS_FALLBACK_COL_ID = 1
PPTS_FALLBACK_COL_NAME = 2
PPTS_FALLBACK_COL_VENDOR = 3
JOURNAL_SHEET_NAME = "Vulnarabilities Publication"


def _clean_header(text: str) -> str:
    return _header_label(text).lower()


def _header_label(text: str) -> str:
    text = re.sub(r"<[^>]+>", " ", str(text))
    text = re.sub(r"\(столбец\s+[A-Z]+\)", "", text, flags=re.IGNORECASE)
    return re.sub(r"\s+", " ", text).strip()


def _find_column(headers: list[str], candidates: tuple[str, ...]) -> int | None:
    for idx, raw_header in enumerate(headers):
        cleaned = _clean_header(raw_header)
        for key in candidates:
            if key in cleaned:
                return idx
    return None


def _find_exact_column(headers: list[str], name: str) -> int | None:
    for idx, raw_header in enumerate(headers):
        if _header_label(raw_header) == name:
            return idx
    return None


def detect_ppts_header_row(rows: list[list[str]]) -> int:
    if not rows:
        return 0

    def cell_a(index: int) -> str:
        if index >= len(rows) or not rows[index]:
            return ""
        return str(rows[index][0]).strip()

    for index in (0, 2):
        value = cell_a(index)
        if value == PPTS_NR_MARK or value.startswith(PPTS_NR_MARK):
            return index
    for index, row in enumerate(rows[:10]):
        value = str(row[0]).strip() if row else ""
        if value == PPTS_NR_MARK or value.startswith(PPTS_NR_MARK):
            return index
    return 0


def detect_ppts_columns(headers: list[str], allow_fallback: bool) -> tuple[int | None, int | None, int | None]:
    col_id = _find_exact_column(headers, PPTS_EXACT_ID)
    col_name = _find_exact_column(headers, PPTS_EXACT_NAME)
    col_vendor = _find_exact_column(headers, PPTS_EXACT_VENDOR)
    if allow_fallback:
        if col_id is None:
            col_id = PPTS_FALLBACK_COL_ID if PPTS_FALLBACK_COL_ID < len(headers) else None
        if col_name is None:
            col_name = PPTS_FALLBACK_COL_NAME if PPTS_FALLBACK_COL_NAME < len(headers) else None
        if col_vendor is None:
            col_vendor = PPTS_FALLBACK_COL_VENDOR if PPTS_FALLBACK_COL_VENDOR < len(headers) else None
    return col_id, col_name, col_vendor


def _cell_str(cell) -> str:
    if cell is None:
        return ""
    return str(cell).strip()


def _read_excel_rows(
    path: Path,
    sheet_name: str | None = None,
    require_sheet: bool = False,
) -> list[list[str]]:
    wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    try:
        if sheet_name and sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
        elif sheet_name and require_sheet:
            raise ReaderError(
                f"Лист {sheet_name!r} не найден в {path.name}. "
                f"Доступные листы: {', '.join(wb.sheetnames)}"
            )
        else:
            ws = wb.active
        if ws is None:
            raise ReaderError(f"No sheet found in {path.name}")
        rows: list[list[str]] = []
        for row in ws.iter_rows(values_only=True):
            rows.append([_cell_str(cell) for cell in row])
        return rows
    finally:
        wb.close()


def _read_csv_rows(path: Path) -> list[list[str]]:
    with open(path, encoding="utf-8-sig", newline="") as f:
        sample = f.read(4096)
        f.seek(0)
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
        except csv.Error:
            dialect = "excel"
        reader = csv.reader(f, dialect)
        return [row for row in reader]


def _read_rows(
    path: Path,
    sheet_name: str | None = None,
    require_sheet: bool = False,
) -> list[list[str]]:
    suffix = path.suffix.lower()
    if suffix == ".xls":
        raise ReaderError(
            "Формат .xls не поддерживается. Сохраните файл как .xlsx или .csv."
        )
    if suffix == ".xlsx":
        return _read_excel_rows(path, sheet_name=sheet_name, require_sheet=require_sheet)
    if suffix == ".csv":
        return _read_csv_rows(path)
    raise ReaderError(f"Unsupported file format: {suffix} (expected .xlsx, .csv)")


def _safe_get(row: list[str], idx: int | None) -> str:
    if idx is None or idx >= len(row):
        return ""
    return row[idx].strip()


_CVE_RE = re.compile(r"CVE-\d{4}-\d+", re.IGNORECASE)


def split_cve_cell(value: str) -> list[str]:
    found = list(dict.fromkeys(m.upper() for m in _CVE_RE.findall(value or "")))
    return found or [value]


_VENDOR_PRODUCT_SEPARATORS = re.compile(r"\s*(?:,\s+|\s+-\s+)\s*")


def _split_vendor_product(combined: str) -> tuple[str, str]:
    if not combined:
        return "", ""

    parts = _VENDOR_PRODUCT_SEPARATORS.split(combined, maxsplit=1)
    if len(parts) == 2:
        return parts[0].strip(), parts[1].strip()
    return "", combined.strip()


TSU_CVE_KEYS = ("cve",)
TSU_CVSS_KEYS = ("cvss",)
TSU_PRODUCT_KEYS = ("продукт", "product", "по", "название")
TSU_SOURCE_KEYS = ("источник", "source", "ссылка", "url")

PPTS_ID_KEYS = ("id пптс", "id", "идентификатор")
PPTS_NAME_KEYS = ("название птс", "название", "наименование", "name", "software")
PPTS_VENDOR_KEYS = ("вендор", "vendor", "производитель", "разработчик")

JOURNAL_COL_RESPONSIBLE = 2
JOURNAL_COL_STATUS = 4
JOURNAL_COL_PPTS_ID = 5
JOURNAL_COL_CVE = 6
JOURNAL_COL_PRODUCT = 8

JOURNAL_CVE_KEYS = ("cve",)
JOURNAL_STATUS_KEYS = ("статус",)
JOURNAL_PPTS_ID_KEYS = ("id пптс", "пптс")
JOURNAL_RESPONSIBLE_KEYS = ("ответственный",)
JOURNAL_PRODUCT_KEYS = ("продукт", "product")


def read_ppts_headers(path: str | Path) -> list[str]:
    path = Path(path)
    if not path.exists():
        raise ReaderError(f"File not found: {path}")
    rows = _read_rows(path)
    if not rows:
        raise ReaderError(f"File {path.name} is empty")
    header_row = detect_ppts_header_row(rows)
    return [_header_label(str(h)) for h in rows[header_row]]


def auto_detect_ppts_mapping(path: str | Path) -> PptsColumnMapping:
    path = Path(path)
    rows = _read_rows(path)
    if not rows:
        raise ReaderError(f"File {path.name} is empty")

    header_row = detect_ppts_header_row(rows)
    raw_headers = [str(h) for h in rows[header_row]]
    found_nr = _header_label(raw_headers[0] if raw_headers else "") == PPTS_NR_MARK or (
        raw_headers and str(raw_headers[0]).strip().startswith(PPTS_NR_MARK)
    )
    col_id, col_name, col_vendor = detect_ppts_columns(raw_headers, allow_fallback=found_nr)
    return PptsColumnMapping(
        file_path=str(path),
        col_id=col_id,
        col_name=col_name,
        col_vendor=col_vendor,
        headers=[_header_label(h) for h in raw_headers],
        header_row=header_row,
    )
